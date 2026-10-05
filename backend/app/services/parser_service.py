import re
import logging
from typing import Dict, Any, List
import pdfplumber
import docx
from app.intelligence.skill_catalog import ALIASES

logger = logging.getLogger(__name__)

TECHNICAL_SKILL_KEYWORDS = [
    "python", "javascript", "typescript", "react", "next.js", "node.js", "fastapi", "flask", "django",
    "sql", "postgresql", "mysql", "mongodb", "redis", "docker", "kubernetes", "aws", "gcp", "azure",
    "machine learning", "deep learning", "tensorflow", "pytorch", "scikit-learn", "pandas", "numpy",
    "data science", "data analysis", "nlp", "computer vision", "generative ai", "langchain", "rag",
    "statistics", "excel", "tableau", "mlops", "llm", "opencv", "transformers", "terraform", "vector databases",
    "git", "ci/cd", "linux", "rest api", "graphql", "c++", "java", "go", "rust", "tailwind css", "html", "css"
]

SOFT_SKILL_KEYWORDS = [
    "leadership", "communication", "problem solving", "critical thinking", "teamwork",
    "collaboration", "adaptability", "time management", "agile", "scrum", "project management",
    "analytical skills", "creativity", "decision making"
]


class ResumeParserService:
    @staticmethod
    def extract_raw_text(file_path: str, file_type: str) -> str:
        try:
            if file_type.lower() == ".pdf":
                with pdfplumber.open(file_path) as pdf:
                    text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            elif file_type.lower() == ".docx":
                document = docx.Document(file_path)
                parts = [p.text for p in document.paragraphs]
                parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                text = "\n".join(parts)
            else:
                raise ValueError("Upload a PDF or DOCX resume.")
        except Exception as exc:
            raise ValueError("Unable to read this resume. Upload a valid PDF or DOCX file.") from exc
        if not text.strip():
            raise ValueError("No readable text found. Use a text-based PDF or DOCX; scanned PDFs need OCR first.")
        return text.strip()

    @staticmethod
    def parse_resume_content(raw_text: str) -> Dict[str, Any]:
        text = raw_text.casefold()
        found = []
        for keyword in TECHNICAL_SKILL_KEYWORDS:
            canonical = next((name for name in ALIASES if name.casefold() == keyword), keyword.title())
            terms = [keyword, *ALIASES.get(canonical, [])]
            if any(re.search(r"(?<!\w)" + re.escape(term) + r"(?!\w)", text) for term in terms):
                found.append(canonical)
        soft = [kw.title() for kw in SOFT_SKILL_KEYWORDS if re.search(r"(?<!\w)" + re.escape(kw) + r"(?!\w)", text)]
        emails = re.findall(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", raw_text)
        contact = {"email": emails[0] if emails else None}
        return {
            "contact": contact,
            "contact_info": contact,
            "technical_skills": found,
            "soft_skills": soft,
            "tools_frameworks": [],
            "education": [],
            "experience": [],
            "projects": [],
        }
