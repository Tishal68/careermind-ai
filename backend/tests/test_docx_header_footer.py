import io
import docx
import pytest
from app.services.parser_service import ResumeParserService
from app.services.upload_safety import extract_isolated


def test_docx_header_and_footer_extraction_survives(tmp_path):
    doc = docx.Document()
    
    # Header with name and email
    header = doc.sections[0].header
    header.paragraphs[0].text = "Alex Sample alex@example.com"
    tbl_h = header.add_table(1, 2, docx.shared.Inches(6))
    tbl_h.cell(0, 0).text = "Phone: +1-555-0199"
    tbl_h.cell(0, 1).text = "Location: Chicago, IL"
    
    # Body with experience and skills
    doc.add_paragraph("Summary: Experienced Python and FastAPI developer.")
    doc.add_paragraph("Skills: Docker, PostgreSQL, Kubernetes, Git.")
    tbl_b = doc.add_table(1, 2)
    tbl_b.cell(0, 0).text = "Senior Software Engineer"
    tbl_b.cell(0, 1).text = "Tech Corp (2020 - Present)"

    # Footer
    footer = doc.sections[0].footer
    footer.paragraphs[0].text = "Certifications: AWS Certified Solutions Architect"
    
    path = tmp_path / "resume_with_header_footer.docx"
    doc.save(path)
    
    # 1. extract_raw_text directly
    raw_text = ResumeParserService.extract_raw_text(str(path), ".docx")
    assert "Alex Sample" in raw_text
    assert "alex@example.com" in raw_text
    assert "Phone: +1-555-0199" in raw_text
    assert "Location: Chicago, IL" in raw_text
    assert "Experienced Python and FastAPI developer." in raw_text
    assert "Senior Software Engineer" in raw_text
    assert "AWS Certified Solutions Architect" in raw_text

    parsed = ResumeParserService.parse_resume_content(raw_text)
    assert parsed["contact"]["email"] == "alex@example.com"
    assert "Python" in parsed["technical_skills"]
    assert "Fastapi" in parsed["technical_skills"]
    assert "Docker" in parsed["technical_skills"]
    assert "Kubernetes" in parsed["technical_skills"]

    # 3. extract_isolated subprocess verification
    isolated_text = extract_isolated(str(path), ".docx")
    assert "alex@example.com" in isolated_text
    assert "AWS Certified Solutions Architect" in isolated_text


def test_docx_linked_sections_do_not_duplicate_header_or_footer(tmp_path):
    doc = docx.Document()
    
    s1 = doc.sections[0]
    s1.header.paragraphs[0].text = "Candidate Header alex.sample@example.com"
    s1.footer.paragraphs[0].text = "Page Footer Note"
    
    doc.add_paragraph("Section 1 Content")
    
    # Section 2 linked to Section 1
    s2 = doc.add_section()
    doc.add_paragraph("Section 2 Content")
    
    # Section 3 linked to Section 2
    s3 = doc.add_section()
    doc.add_paragraph("Section 3 Content")
    
    path = tmp_path / "multisection_linked.docx"
    doc.save(path)
    
    raw_text = ResumeParserService.extract_raw_text(str(path), ".docx")
    assert raw_text.count("Candidate Header alex.sample@example.com") == 1
    assert raw_text.count("Page Footer Note") == 1
    assert "Section 1 Content" in raw_text
    assert "Section 2 Content" in raw_text
    assert "Section 3 Content" in raw_text


def test_docx_unlinked_sections_extract_all_unique_headers(tmp_path):
    doc = docx.Document()
    
    s1 = doc.sections[0]
    s1.header.paragraphs[0].text = "Header Section 1"
    doc.add_paragraph("Content 1")
    
    s2 = doc.add_section()
    s2.header.is_linked_to_previous = False
    s2.header.paragraphs[0].text = "Header Section 2"
    doc.add_paragraph("Content 2")
    
    path = tmp_path / "multisection_unlinked.docx"
    doc.save(path)
    
    raw_text = ResumeParserService.extract_raw_text(str(path), ".docx")
    assert "Header Section 1" in raw_text
    assert "Header Section 2" in raw_text
