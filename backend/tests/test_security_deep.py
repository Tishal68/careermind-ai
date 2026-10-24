import io
import json
import zipfile
from unittest.mock import Mock
from datetime import datetime, timezone
import httpx
import pytest
from fastapi import FastAPI, HTTPException, Request
from fastapi.testclient import TestClient
from pydantic import SecretStr
from test_security_regressions import secured, account
from app.core.request_limits import RequestLimitsMiddleware, WindowLimiter
from app.core.config import Settings, settings
from app.core.security import create_access_token
from app.services.upload_safety import preflight_docx, extract_isolated, remove_owned_directory
from app.services.job_research import public_url
from app.services.interview_service import InterviewService, ai_router_engine
from app.models.interview import InterviewSession
from app.models.user import User
from app.ai.providers.ollama import OllamaLLMProvider


def test_logout_revokes_only_the_signed_out_token(secured):
    client, db = secured
    uid, headers = account(client)
    second = {"Authorization": "Bearer " + create_access_token(uid, session_version=db.get(User, uid).session_version)}
    assert client.post("/api/v1/auth/logout", headers=headers).status_code == 204
    assert client.get("/api/v1/auth/me", headers=headers).status_code == 401
    assert client.get("/api/v1/auth/me", headers=second).status_code == 200


def test_request_limit_counts_actual_bytes_and_auth_attempts():
    app = FastAPI()
    app.add_middleware(RequestLimitsMiddleware, auth_limit=2)
    @app.post('/api/v1/auth/login')
    def login():
        return {}
    @app.post('/data')
    async def data(request: Request):
        return {"size": len(await request.body())}
    with TestClient(app) as client:
        assert client.post('/data', content=b'x' * 65536).status_code == 200
        assert client.post('/data', content=b'x' * 65537).status_code == 413
        assert client.post('/data', content=iter([b'x'*40000, b'y'*40000])).status_code == 413
        assert client.post('/api/v1/auth/login').status_code == 200
        assert client.post('/api/v1/auth/login').status_code == 200
        response = client.post('/api/v1/auth/login', headers={"X-Forwarded-For": "new-ip"})
        assert response.status_code == 429
        assert response.headers['retry-after'] == '60'


def test_limiter_is_bounded_and_expires(monkeypatch):
    import app.core.request_limits as limits
    now = [1.0]
    monkeypatch.setattr(limits.time, 'monotonic', lambda: now[0])
    limiter = WindowLimiter(limit=2, window=10, max_keys=1)
    assert limiter.allow('a') and limiter.allow('a')
    assert not limiter.allow('a') and not limiter.allow('b')
    now[0] = 12.0
    assert limiter.allow('b')


def test_docx_decompression_bomb_rejected(tmp_path):
    path = tmp_path/'bomb.docx'
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', b'x' * (2 * 1024 * 1024))
    with pytest.raises(ValueError):
        preflight_docx(path)


def test_document_timeout_is_actionable(monkeypatch):
    import subprocess
    import app.services.upload_safety as safety
    monkeypatch.setattr(safety.subprocess, 'run', Mock(side_effect=subprocess.TimeoutExpired('parser',20)))
    with pytest.raises(ValueError, match='took too long'):
        extract_isolated('unused.pdf', '.pdf')


def test_deletion_cannot_escape_upload_root(tmp_path):
    root = tmp_path/'uploads';root.mkdir()
    outside = tmp_path/'important';outside.mkdir()
    (outside/'keep.txt').write_text('keep')
    with pytest.raises(HTTPException):
        remove_owned_directory(outside,root)
    with pytest.raises(HTTPException):
        remove_owned_directory(root,root)
    assert (outside/'keep.txt').read_text() == 'keep'


@pytest.mark.parametrize('url', ['https://127.0.0.1./x','https://127.1/x','https://0177.0.0.1/x','https://x.internal./x','https://example.com\\@127.0.0.1/x','https://%31%32%37.0.0.1/x','https://example.com/\nprivate'])
def test_ambiguous_job_urls_are_rejected(url):
    with pytest.raises(HTTPException):
        public_url(url)


def test_cloud_secret_not_forwarded_to_custom_endpoint(monkeypatch):
    monkeypatch.setattr(settings, 'OLLAMA_API_KEY', SecretStr('private-cloud-key'))
    provider = OllamaLLMProvider(base_url='http://localhost:11434')
    assert provider.api_key == ''
    custom = OllamaLLMProvider(base_url='https://private.example.com',api_key='explicit-private-key')
    assert custom.api_key == 'explicit-private-key'


def test_interview_provider_failure_does_not_create_fake_session(secured,monkeypatch):
    client, db = secured
    _, headers = account(client)
    monkeypatch.setattr(ai_router_engine,'generate_text',Mock(side_effect=HTTPException(503,'Provider unavailable')))
    response=client.post('/api/v1/interview/start',headers=headers,json={})
    assert response.status_code==503
    assert db.query(InterviewSession).count()==0


def test_interview_ownership_persistence_and_replay(secured,monkeypatch):
    client, db=secured
    _,headers=account(client)
    _,foreign=account(client,'foreign@example.com')
    monkeypatch.setattr(ai_router_engine,'generate_text',Mock(return_value='Explain Python generators.'))
    feedback={'score':40,'strengths':['Mentions iteration'],'mistakes':['Explain lazy evaluation'],'suggested_answer':'Generators yield items lazily.','next_question':'Explain yield.'}
    monkeypatch.setattr(ai_router_engine,'generate_json',Mock(return_value=feedback))
    session=client.post('/api/v1/interview/start',headers=headers,json={}).json()
    body={'session_id':session['id'],'question':session['history'][0]['question'],'user_answer':'An iterator.'}
    assert client.post('/api/v1/interview/evaluate',headers=foreign,json=body).status_code==404
    assert client.post('/api/v1/interview/evaluate',headers=headers,json=body).status_code==200
    db.expire_all()
    saved=db.get(InterviewSession,session['id'])
    assert saved.overall_score==40 and saved.history[0]['answer']=='An iterator.'
    assert client.post('/api/v1/interview/evaluate',headers=headers,json=body).status_code==200
    assert ai_router_engine.generate_json.call_count == 1
    assert client.post('/api/v1/interview/evaluate',headers=headers,json={**body,'user_answer':'Different answer'}).status_code==409


@pytest.mark.parametrize('feedback',[{}, {'score':999}, {'score':float('nan')}])
def test_invalid_interview_feedback_is_not_scored(monkeypatch,feedback):
    monkeypatch.setattr(ai_router_engine,'generate_json',Mock(return_value=feedback))
    with pytest.raises(HTTPException) as error:
        InterviewService.evaluate_response('Technical','Engineer','Question','Answer')
    assert error.value.status_code==502


def test_admin_deletion_removes_private_files(secured,tmp_path,monkeypatch):
    client,db=secured
    uid,headers=account(client)
    other,_=account(client,'delete@example.com')
    db.get(User,uid).is_superuser=True;db.commit()
    monkeypatch.setattr(settings,'UPLOAD_DIR',str(tmp_path))
    directory=tmp_path/f'user_{other}'
    directory.mkdir();(directory/'resume.pdf').write_bytes(b'private')
    assert client.delete(f'/api/v1/admin/users/{other}',headers=headers).status_code==204
    assert not directory.exists()


def test_production_requires_stable_secret(monkeypatch):
    monkeypatch.delenv('SECRET_KEY',raising=False)
    with pytest.raises(ValueError):
        Settings(ENVIRONMENT='production',_env_file=None)
    assert Settings(ENVIRONMENT='production',SECRET_KEY='x'*48,_env_file=None)


def test_new_account_has_no_fabricated_progress(secured):
    client,_=secured
    _,headers=account(client)
    data=client.get('/api/v1/dashboard/overview',headers=headers).json()
    assert data['resumes_count']==data['completed_projects_count']==0
    assert data['learning_progress_percent'] is None


def test_api_responses_are_not_cacheable(application_client):
    response=application_client.get('/api/v1/auth/me')
    assert response.status_code==401
    assert response.headers['cache-control']=='no-store'
    assert response.headers['x-content-type-options']=='nosniff'


def test_document_worker_reads_docx_outside_backend(tmp_path, monkeypatch):
    from docx import Document
    document = Document()
    document.add_paragraph("Resume: Python developer with SQL experience.")
    document.save(tmp_path / "resume.docx")
    monkeypatch.chdir(tmp_path)
    assert "Python developer" in extract_isolated("resume.docx", ".docx")


def test_document_worker_reports_empty_document(tmp_path, monkeypatch):
    from docx import Document
    Document().save(tmp_path / "empty.docx")
    monkeypatch.chdir(tmp_path)
    with pytest.raises(ValueError, match="No readable text"):
        extract_isolated("empty.docx", ".docx")


def test_document_worker_crash_is_service_error(monkeypatch):
    import app.services.upload_safety as safety
    monkeypatch.setattr(safety.subprocess, 'run', Mock(return_value=Mock(returncode=1, stdout='')))
    with pytest.raises(HTTPException) as exc:
        extract_isolated('resume.pdf', '.pdf')
    assert exc.value.status_code == 503
